import re
from typing import Tuple, List

import openpyxl
from openpyxl.cell import Cell
from openpyxl.styles import PatternFill
from openpyxl.styles.colors import Color
from openpyxl.workbook import Workbook

GROUP_COLOR = "FFFFFF00"
REMOTED_COLOR = "FFD9D9D9"


def get_group_name(string: str):
    if not isinstance(string, str):
        return None

    match = re.search(
        r"(\d?[А-яЕЁ]+-\d.\d\dК?)",
        string.upper().strip(),
        re.RegexFlag.MULTILINE | re.RegexFlag.IGNORECASE
    )

    if match:
        return match.group()


def is_not_empty_text(string: str):
    return bool(re.search(
        r"[А-яЕЁ\d]",
        string,
        re.RegexFlag.MULTILINE | re.RegexFlag.IGNORECASE
    )) if string else False


def get_groups(cells: Tuple[Cell]):
    groups = {}
    for index, cell in enumerate(cells):
        text = cell.value

        fill: PatternFill = cell.fill
        fg_color: Color = fill.fgColor
        color = fg_color.rgb

        if not text or color != GROUP_COLOR:
            continue

        group_name = get_group_name(text)
        if not group_name:
            continue

        groups.update({index: group_name})

    return groups


def get_data_groups(data: List[Tuple]):
    data_groups = {}
    for row_index, cells in enumerate(data):
        groups = get_groups(cells)
        if not groups:
            continue

        data_groups[row_index] = {
            "groups": list(groups.values()),
            "columns": list(groups.keys())
        }

    return data_groups


def get_columns_groups(data_groups: dict):
    name_groups = []
    for value in data_groups.values():
        name_groups = [*name_groups, *value["groups"]]

    return {
        str(index): name for index, name in enumerate(name_groups)
    }


def get_numbers(data: List[Tuple[Cell]], data_groups: dict):
    data_group = None
    numbers = {"normal": [], "indexes": {}}
    for row_index, cells in enumerate(data):
        if data_groups.get(row_index):
            data_group = data_groups[row_index]
            continue
        if not data_group:
            continue

        cell_index = data_group["columns"][0] - 1

        value = cells[cell_index].value
        if not value:
            value = numbers["normal"][-1]
        else:
            numbers["normal"].append(value)

        indexes_on_value = numbers["indexes"].get(numbers["normal"].index(value, -1), [])
        indexes_on_value.append(row_index)

        numbers["indexes"][numbers["normal"].index(value, -1)] = indexes_on_value

    return numbers


def get_raw_schedule(data: List[Tuple], data_groups: dict, number_indexes: dict):
    data_group = None
    raw_schedule = {}
    for row_index, cells in enumerate(data):
        if data_groups.get(row_index):
            data_group = data_groups[row_index]
            continue
        if not data_group:
            continue

        groups = data_group["groups"]
        indexes = data_group["columns"]

        for index, group in zip(indexes, groups):
            cell: Cell = cells[index]

            text: str = cell.value.strip() if is_not_empty_text(cell.value) else None

            fill: PatternFill = cell.fill
            fg_color: Color = fill.fgColor
            color = fg_color.rgb

            remoted = color == REMOTED_COLOR

            group_schedule = raw_schedule.get(group, [])

            for num_index, row_indexes in number_indexes.items():
                if row_index not in row_indexes:
                    continue

                try:
                    tape: dict = group_schedule[num_index]
                except IndexError:
                    group_schedule.append({"text": text, "remoted": remoted})
                else:
                    if not text:
                        break

                    tape.update(text=f"{tape['text']}\n{text}")

                    group_schedule[num_index] = tape

                raw_schedule[group] = group_schedule
                break

    for group, schedule in raw_schedule.items():
        group_schedule = schedule.copy()

        for tape in reversed(schedule):
            if not tape["text"]:
                group_schedule.pop()
            else:
                break

        raw_schedule[group] = group_schedule

    return raw_schedule


def parse_schedule(file_path: str):
    wb: Workbook = openpyxl.load_workbook(file_path)
    ws = wb.worksheets[0]

    data = [row for row in ws]
    wb.close()

    result = {}

    data_groups = get_data_groups(data)

    result["groups"] = get_columns_groups(data_groups)

    numbers = get_numbers(data, data_groups)
    result["numbers"] = numbers["normal"]

    raw_schedule = get_raw_schedule(data, data_groups, numbers["indexes"])

    result["schedule"] = {
        str(index): tapes for index, tapes in
        zip(result["groups"], raw_schedule.values())
    }

    return result
