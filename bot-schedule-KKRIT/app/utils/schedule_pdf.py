import re
import os
from typing import Tuple, List

from docx import Document
from docx.document import Document as FileDocument
from docx.oxml import CT_Tc, CT_TcPr, CT_Tbl, CT_Row, CT_P, CT_R, CT_Text
from docx.table import _Cell, _Rows
from lxml.etree import _Element

from app.utils.sync import run_async

STANDARD_COLOR = "ffffff"
GROUP_COLOR = "ffff00"
REMOTED_COLOR = "e0e0e0"

FILL_KEY = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill"


def get_shd_cell(tc: CT_Tc):
    for tcPr in tc:
        if not isinstance(tcPr, CT_TcPr):
            continue

        for element in tcPr:
            if not isinstance(element, _Element) or element.tag[-3:] != "shd":
                continue

            return element.attrib[FILL_KEY]


def get_shd_cell_text(tc: CT_Tc):
    # <w:tbl> -> <w:tr> -> <w:tc> -> <w:tcPr> -> <w:shd>

    for tcPr in tc:
        if not isinstance(tcPr, CT_Tbl):
            continue

        for elements_tr_pr in tcPr:
            if not isinstance(elements_tr_pr, CT_Row) and elements_tr_pr.tag[-3:] != "tr":
                continue
            for elements_tc in elements_tr_pr:
                if not isinstance(elements_tc, CT_Tc):
                    continue
                for elements_tr_pr_tbl in elements_tc:
                    if not isinstance(elements_tr_pr_tbl, CT_TcPr):
                        continue

                    for element in elements_tr_pr_tbl:
                        if not isinstance(element, _Element) or element.tag[-3:] != "shd":
                            continue

                        return element.attrib[FILL_KEY]


def get_shd(tc: CT_Tc):
    fill = get_shd_cell(tc)
    if fill:
        fill = str(fill).lower()
        if fill in (GROUP_COLOR, REMOTED_COLOR):
            return fill

    text_fill = get_shd_cell_text(tc)
    return str(text_fill).lower() if text_fill else None


def get_group_name(string: str):
    if not isinstance(string, str):
        return None

    match = re.search(
        r"(\d?[А-яЕЁ]+-\d.\d\dК?)",
        string.upper().strip(),
        re.RegexFlag.MULTILINE | re.RegexFlag.IGNORECASE
    )

    if match:
        return match.group()


def get_groups(cells: Tuple[_Cell]):
    groups = {}
    for index, cell in enumerate(cells):
        tc: CT_Tc = cell._tc

        text = get_cell_text(cell)
        shd = get_shd(tc)

        if not text or shd != GROUP_COLOR:
            continue

        group_name = get_group_name(text)
        if not group_name:
            continue

        groups.update({index: group_name})

    return groups


def check_empty_text(string: str):
    return not re.search(
        r"[А-яЕЁ\d]",
        string,
        re.RegexFlag.MULTILINE | re.RegexFlag.IGNORECASE
    )


def get_data_groups(data: List[Tuple]):
    data_groups = {}
    for row_index, cells in enumerate(data):
        groups = get_groups(cells)
        if not groups:
            continue

        data_groups[row_index] = {
            "groups": list(groups.values()),
            "indexes": list(groups.keys()),
            "start_cell_index": list(groups.keys())[0]
        }

    return data_groups


def get_indexes_groups(data_groups: dict):
    name_groups = []
    for value in data_groups.values():
        name_groups = [*name_groups, *value["groups"]]

    return {
        str(index): name for index, name in enumerate(name_groups)
    }


def get_numbers(data: List[Tuple], data_groups: dict):
    data_group = None
    numbers = []
    for row_index, cells in enumerate(data):
        if data_groups.get(row_index):
            data_group = data_groups[row_index]
            continue

        cell_index = data_group["start_cell_index"] - 1
        numbers.append(cells[cell_index].text)

    return numbers


def get_cell_text(cell: _Cell):
    ps = []
    tbls = []
    for element in cell._tc:
        if isinstance(element, CT_P):
            ps.append(element)
        if isinstance(element, CT_Tbl):
            tbls.append(element)

    text = []
    for p in ps:
        for r in p:
            if not isinstance(r, CT_R):
                continue
            for t in r:
                if not isinstance(t, CT_Text):
                    continue

                text.append(t.text)

    for tbl in tbls:
        for tr in tbl:
            if not isinstance(tr, CT_Row) and tr.tag[-3:] != "tr":
                continue
            for tc in tr:
                if not isinstance(tc, CT_Tc):
                    continue
                for p in tc:
                    if not isinstance(p, CT_P):
                        continue
                    for r in p:
                        if not isinstance(r, CT_R):
                            continue
                        for t in r:
                            if not isinstance(t, CT_Text):
                                continue

                            text.append(t.text)

    row_text = " ".join(text)

    return row_text


def get_raw_schedule(data: List[Tuple], data_groups: dict):
    data_group = None
    raw_schedule = {}
    for row_index, cells in enumerate(data):
        if data_groups.get(row_index):
            data_group = data_groups[row_index]
            continue

        groups = data_group["groups"]
        indexes = data_group["indexes"]

        for index, group in zip(indexes, groups):
            cell: _Cell = cells[index]

            row_text = get_cell_text(cell)
            text = row_text if not check_empty_text(row_text) else None

            shd = get_shd(cell._tc)

            remoted = shd == REMOTED_COLOR

            # TODO: Найти когда нибудь нормальный способ что бы определять серый фон (до), а то сейчас хуй определить
            if text:
                split_text = text.lower().split()
                try:
                    remoted_index = split_text.index("до")
                except ValueError:
                    pass
                else:
                    remoted = True

                    text = text.split()
                    del text[remoted_index]

                    text = " ".join(text)

            group_schedule = raw_schedule.get(group, [])
            group_schedule.append({"text": text, "remoted": remoted})

            raw_schedule[group] = group_schedule

    for group, schedule in raw_schedule.items():
        group_schedule = schedule.copy()

        for tape in reversed(schedule):
            if not tape["text"]:
                group_schedule.pop()
            else:
                break

        raw_schedule[group] = group_schedule

    return raw_schedule


def get_data_schedule(rows: _Rows):
    result = {
        "groups": {},
        "schedule": {},
        "numbers": {}
    }

    data = [row.cells for row in rows]

    data_groups = get_data_groups(data)

    result["groups"] = get_indexes_groups(data_groups)
    result["numbers"] = get_numbers(data, data_groups)

    raw_schedule = get_raw_schedule(data, data_groups)

    result["schedule"] = {
        str(index): tapes for index, tapes in
        zip(result["groups"], raw_schedule.values())
    }

    return result


def parse_schedule(file_path: str):
    from pdf2docx import Converter

    converted_file_path = f"{file_path}.docx"

    # if not os.path.isfile(converted_file_path):
    #     cv = Converter(file_path)
    #     cv.convert(converted_file_path)
    #     cv.close()

    with open(converted_file_path, "rb") as file:
        document: FileDocument = Document(file)
    return document

    # return get_data_schedule(document.tables[-1].rows), converted_file_path


async def go_parse_schedule(file_path: str):
    return await run_async(parse_schedule, file_path)


document = parse_schedule("01.09.2023.pdf")
print(
    [
        [
            get_cell_text(cell) for cell in row.cells
        ]
        for row in document.tables[0].rows
    ]
)


def get_groups_new(cells: Tuple[_Cell]):
    groups = []
    for cell in cells:
        text = get_cell_text(cell)
        if not text and text.lower() == "группы":
            continue

        group_name = get_group_name(text)
        if not group_name:
            continue

        groups.append(group_name)

    return groups


def get_schedule_new(rows: _Rows):
    all_cells: List[List[_Cell]] = [row.cells for row in rows]

    for cells in all_cells:
        ...
